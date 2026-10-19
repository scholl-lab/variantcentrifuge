"""Regression tests for eliminating openpyxl reloads during Excel generation (#159)."""

import json
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd
import pytest

from variantcentrifuge.pipeline_core.context import PipelineContext
from variantcentrifuge.pipeline_core.workspace import Workspace
from variantcentrifuge.stages.output_stages import ExcelReportStage


@pytest.fixture
def excel_context(tmp_path: Path) -> PipelineContext:
    """Create a pipeline context with all optional sheets and IGV report mapping."""
    output_dir = tmp_path / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    report_dir = output_dir / "report" / "igv"
    report_dir.mkdir(parents=True, exist_ok=True)

    # 1. Setup IGV reports map
    igv_map = {
        "variants": [
            {
                "chrom": "chr1",
                "pos": 100,
                "ref": "A",
                "alt": "T",
                "sample_reports": {"Sample1": "igv/Sample1_var1.html"},
            },
            {
                "chrom": "chr2",
                "pos": 200,
                "ref": "C",
                "alt": "G",
                "sample_reports": {
                    "Sample1": "igv/Sample1_var2.html",
                    "Sample2": "igv/Sample2_var2.html",
                },
            },
        ]
    }
    with open(report_dir / "igv_reports_map.json", "w", encoding="utf-8") as f:
        json.dump(igv_map, f)

    # 2. Results DataFrame
    results_df = pd.DataFrame(
        {
            "CHROM": ["chr1", "chr2", "chr3"],
            "POS": [100, 200, 300],
            "REF": ["A", "C", "G"],
            "ALT": ["T", "G", "A"],
            "GT": ["Sample1(0/1)", "Sample1(0/1);Sample2(0/1)", "Sample3(0/1)"],
            "ClinVar": [
                "https://www.ncbi.nlm.nih.gov/clinvar/?term=chr1-100-A-T",
                "https://www.ncbi.nlm.nih.gov/clinvar/?term=chr2-200-C-G",
                "https://www.ncbi.nlm.nih.gov/clinvar/?term=chr3-300-G-A",
            ],
        }
    )

    # 3. Create TSV files for optional sheets
    metadata_tsv = output_dir / "metadata.tsv"
    stats_tsv = output_dir / "statistics.tsv"
    gene_burden_tsv = output_dir / "gene_burden.tsv"
    assoc_tsv = output_dir / "association.tsv"
    tsv_output = output_dir / "variants.tsv"

    results_df.to_csv(tsv_output, sep="\t", index=False)
    pd.DataFrame({"Param": ["version", "build"], "Value": ["1.0", "GRCh38"]}).to_csv(
        metadata_tsv, sep="\t", index=False
    )
    pd.DataFrame({"Metric": ["total", "filtered"], "Count": [10, 2]}).to_csv(
        stats_tsv, sep="\t", index=False
    )
    pd.DataFrame({"Gene": ["GENE1"], "P_value": [0.001]}).to_csv(
        gene_burden_tsv, sep="\t", index=False
    )
    pd.DataFrame({"Gene": ["GENE1"], "Assoc_P": [0.005]}).to_csv(assoc_tsv, sep="\t", index=False)

    workspace = Workspace(output_dir, "variants")
    context = PipelineContext(
        args=None,
        config={
            "xlsx": True,
            "output_file": str(tsv_output),
            "no_stats": False,
            "stats_output_file": str(stats_tsv),
            "perform_gene_burden": True,
            "gene_burden_output": str(gene_burden_tsv),
            "perform_association": True,
            "association_output": str(assoc_tsv),
            "igv_enabled": True,
        },
        workspace=workspace,
    )
    context.final_output_path = tsv_output
    context.current_dataframe = results_df
    context.report_paths["metadata"] = metadata_tsv
    context.mark_complete("dataframe_loading")
    context.mark_complete("tsv_output")
    context.mark_complete("metadata_generation")
    context.mark_complete("statistics_generation")

    return context


@pytest.mark.unit
def test_excel_report_stage_performs_no_openpyxl_reload(excel_context: PipelineContext):
    """ExcelReportStage writes Results and optional sheets without loading via openpyxl."""
    stage = ExcelReportStage()

    with (
        patch("openpyxl.load_workbook", wraps=openpyxl.load_workbook) as spy_load,
        patch(
            "variantcentrifuge.converter.load_workbook", wraps=openpyxl.load_workbook
        ) as spy_load_conv,
    ):
        stage(excel_context)
        assert spy_load.call_count == 0, f"Expected 0 openpyxl calls, got {spy_load.call_count}"
        assert spy_load_conv.call_count == 0, (
            f"Expected 0 converter load_workbook calls, got {spy_load_conv.call_count}"
        )

    excel_file = excel_context.workspace.output_dir / "variants.xlsx"
    assert excel_file.exists()

    # Now verify the produced file using openpyxl in test verification
    wb = openpyxl.load_workbook(excel_file)
    assert wb.sheetnames == ["Results", "Metadata", "Statistics", "Gene Burden", "Association"]

    for name in wb.sheetnames:
        ws = wb[name]
        assert ws.freeze_panes == "A2", f"Sheet {name} missing freeze_panes A2"
        assert ws.auto_filter.ref is not None, f"Sheet {name} missing autofilter"

    # Verify Results sheet contents, hyperlinks, and IGV Report Links
    results_ws = wb["Results"]
    headers = [cell.value for cell in results_ws[1]]
    assert "ClinVar" in headers
    assert "IGV Report Links" in headers

    clinvar_col = headers.index("ClinVar") + 1
    igv_col = headers.index("IGV Report Links") + 1

    # Row 2 (first variant): ClinVar hyperlink and single IGV link
    cell_clinvar_1 = results_ws.cell(row=2, column=clinvar_col)
    assert cell_clinvar_1.value == "ClinVar"
    expected_clinvar_url = "https://www.ncbi.nlm.nih.gov/clinvar/?term=chr1-100-A-T"
    assert expected_clinvar_url in cell_clinvar_1.hyperlink.target

    cell_igv_1 = results_ws.cell(row=2, column=igv_col)
    assert cell_igv_1.value == "Sample1"
    assert cell_igv_1.hyperlink.target.replace("\\", "/") == "report/igv/Sample1_var1.html"

    # Row 3 (second variant): multiple IGV reports with comment
    cell_igv_2 = results_ws.cell(row=3, column=igv_col)
    assert cell_igv_2.value == "Sample1 (+1 others)"
    assert cell_igv_2.comment is not None
    assert "Sample1 (report/igv/Sample1_var2.html)" in cell_igv_2.comment.text
    assert "Sample2 (report/igv/Sample2_var2.html)" in cell_igv_2.comment.text

    # Row 4 (third variant): no IGV report mapped -> N/A
    cell_igv_3 = results_ws.cell(row=4, column=igv_col)
    assert cell_igv_3.value == "N/A"

    # Verify optional sheets cell values
    meta_ws = wb["Metadata"]
    assert [c.value for c in meta_ws[1]] == ["Param", "Value"]
    assert [c.value for c in meta_ws[2]] == ["version", "1.0"]
    assert [c.value for c in meta_ws[3]] == ["build", "GRCh38"]

    stats_ws = wb["Statistics"]
    assert [c.value for c in stats_ws[1]] == ["Metric", "Count"]
    assert [c.value for c in stats_ws[2]] == ["total", 10]
    assert [c.value for c in stats_ws[3]] == ["filtered", 2]

    burden_ws = wb["Gene Burden"]
    assert [c.value for c in burden_ws[1]] == ["Gene", "P_value"]
    assert [c.value for c in burden_ws[2]] == ["GENE1", 0.001]

    assoc_ws = wb["Association"]
    assert [c.value for c in assoc_ws[1]] == ["Gene", "Assoc_P"]
    assert [c.value for c in assoc_ws[2]] == ["GENE1", 0.005]


@pytest.mark.unit
def test_excel_empty_results_and_association_fidelity(tmp_path: Path):
    """Verify empty Results sheet (0 variants) with Association sheet and formatting."""
    output_dir = tmp_path / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    tsv_output = output_dir / "empty.tsv"
    assoc_tsv = output_dir / "association.tsv"

    empty_df = pd.DataFrame(columns=["CHROM", "POS", "REF", "ALT", "GT"])
    empty_df.to_csv(tsv_output, sep="\t", index=False)
    pd.DataFrame({"Gene": ["TP53"], "Assoc_P": [0.01]}).to_csv(assoc_tsv, sep="\t", index=False)

    workspace = Workspace(output_dir, "variants")
    context = PipelineContext(
        args=None,
        config={
            "xlsx": True,
            "output_file": str(tsv_output),
            "no_stats": True,
            "perform_association": True,
            "association_output": str(assoc_tsv),
        },
        workspace=workspace,
    )
    context.final_output_path = tsv_output
    context.current_dataframe = empty_df
    context.mark_complete("dataframe_loading")
    context.mark_complete("tsv_output")
    context.mark_complete("metadata_generation")
    context.mark_complete("statistics_generation")

    stage = ExcelReportStage()
    with (
        patch("openpyxl.load_workbook", wraps=openpyxl.load_workbook) as spy_load,
        patch(
            "variantcentrifuge.converter.load_workbook", wraps=openpyxl.load_workbook
        ) as spy_load_conv,
    ):
        stage(context)
        assert spy_load.call_count == 0
        assert spy_load_conv.call_count == 0

    excel_file = context.workspace.get_output_path("", ".xlsx")
    assert excel_file.exists()

    wb = openpyxl.load_workbook(excel_file)
    assert wb.sheetnames == ["Results", "Association"]
    assert wb["Results"].max_row == 1  # Header only
    assert wb["Association"].max_row == 2  # Header + 1 row


@pytest.mark.unit
def test_excel_legacy_list_format_igv_map(tmp_path: Path):
    """Verify convert_to_excel supports top-level list IGV maps (#159)."""
    from variantcentrifuge.converter import convert_to_excel

    output_dir = tmp_path / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    report_dir = output_dir / "report" / "igv"
    report_dir.mkdir(parents=True, exist_ok=True)

    # Legacy top-level list format with flat report_path
    legacy_map = [
        {
            "chrom": "chr1",
            "pos": 100,
            "ref": "A",
            "alt": "T",
            "sample_id": "Sample1",
            "report_path": "igv/Sample1_legacy.html",
        }
    ]
    with open(report_dir / "igv_reports_map.json", "w", encoding="utf-8") as f:
        json.dump(legacy_map, f)

    df = pd.DataFrame(
        {
            "CHROM": ["chr1"],
            "POS": [100],
            "REF": ["A"],
            "ALT": ["T"],
            "GT": ["Sample1(0/1)"],
        }
    )
    xlsx_file = output_dir / "variants.xlsx"
    convert_to_excel(
        str(xlsx_file),
        cfg={"igv_enabled": True},
        df=df,
        output_file=str(xlsx_file),
    )

    wb = openpyxl.load_workbook(xlsx_file)
    ws = wb["Results"]
    headers = [cell.value for cell in ws[1]]
    assert "IGV Report Links" in headers
    igv_cell = ws.cell(row=2, column=headers.index("IGV Report Links") + 1)
    assert igv_cell.value == "Sample1"
    assert "igv/Sample1_legacy.html" in igv_cell.hyperlink.target.replace("\\", "/")


@pytest.mark.unit
def test_excel_6000_rows_6_urls_no_hyperlink_limit_exhaustion(tmp_path: Path):
    """Verify 6,000 rows x 6 URL cols + IGV col (42,000 URLs) don't hit 65,530 limit (#159)."""
    import warnings

    from variantcentrifuge.converter import convert_to_excel

    output_dir = tmp_path / "stress_output"
    output_dir.mkdir(parents=True, exist_ok=True)
    report_dir = output_dir / "report" / "igv"
    report_dir.mkdir(parents=True, exist_ok=True)

    n_rows = 6000
    legacy_map = [
        {
            "chrom": "chr1",
            "pos": row,
            "ref": "A",
            "alt": "T",
            "sample_id": "Sample1",
            "report_path": f"igv/Sample1_{row}.html",
        }
        for row in range(1, n_rows + 1)
    ]
    with open(report_dir / "igv_reports_map.json", "w", encoding="utf-8") as f:
        json.dump(legacy_map, f)

    data = {
        "CHROM": ["chr1"] * n_rows,
        "POS": list(range(1, n_rows + 1)),
        "REF": ["A"] * n_rows,
        "ALT": ["T"] * n_rows,
        "GT": ["Sample1(0/1)"] * n_rows,
    }
    for col_idx in range(1, 7):
        data[f"URL_COL_{col_idx}"] = [
            f"https://example.org/path/{col_idx}/{row}" for row in range(n_rows)
        ]

    df = pd.DataFrame(data)
    xlsx_file = output_dir / "large_urls.xlsx"

    # With strings_to_urls=False and single registration, exactly 42,000 hyperlinks are written
    with warnings.catch_warnings(record=True) as recorded:
        warnings.simplefilter("always")
        convert_to_excel(
            str(xlsx_file),
            cfg={"igv_enabled": True},
            df=df,
            output_file=str(xlsx_file),
        )

    limit_warnings = [
        w for w in recorded if "exceeds Excel's limit of 65,530 URLs" in str(w.message)
    ]
    assert len(limit_warnings) == 0

    assert xlsx_file.exists()
    wb = openpyxl.load_workbook(xlsx_file)
    ws = wb["Results"]
    assert ws.max_row == n_rows + 1

    headers = [cell.value for cell in ws[1]]
    assert "IGV Report Links" in headers
    igv_col_idx = headers.index("IGV Report Links") + 1

    # Verify first and last rows have proper hyperlinks
    first_url_cell = ws.cell(row=2, column=6)
    assert first_url_cell.value == "URL_COL_1"
    assert first_url_cell.hyperlink.target == "https://example.org/path/1/0"

    last_url_cell = ws.cell(row=n_rows + 1, column=11)
    assert last_url_cell.value == "URL_COL_6"
    assert last_url_cell.hyperlink.target == f"https://example.org/path/6/{n_rows - 1}"

    first_igv_cell = ws.cell(row=2, column=igv_col_idx)
    assert first_igv_cell.value == "Sample1"
    assert "igv/Sample1_1.html" in first_igv_cell.hyperlink.target.replace("\\", "/")

    last_igv_cell = ws.cell(row=n_rows + 1, column=igv_col_idx)
    assert last_igv_cell.value == "Sample1"
    assert f"igv/Sample1_{n_rows}.html" in last_igv_cell.hyperlink.target.replace("\\", "/")

    # Assert that all 36,000 URL cells and 6,000 IGV link cells have active hyperlinks
    url_hlinks = sum(
        1
        for row in ws.iter_rows(min_row=2, max_row=n_rows + 1, min_col=6, max_col=11)
        for cell in row
        if cell.hyperlink is not None
    )
    assert url_hlinks == 36000

    igv_hlinks = sum(
        1
        for row in ws.iter_rows(
            min_row=2, max_row=n_rows + 1, min_col=igv_col_idx, max_col=igv_col_idx
        )
        for cell in row
        if cell.hyperlink is not None
    )
    assert igv_hlinks == 6000


@pytest.mark.unit
def test_excel_generation_benchmark_performance_and_memory(tmp_path: Path, record_property):
    """Deterministic structural/memory comparison: multi-pass vs single-pass writer (#159)."""
    import gc
    import time
    import tracemalloc

    from openpyxl.utils import get_column_letter

    from variantcentrifuge.converter import convert_to_excel

    n_rows = 2000
    df = pd.DataFrame(
        {
            "CHROM": ["chr1"] * n_rows,
            "POS": list(range(n_rows)),
            "REF": ["A"] * n_rows,
            "ALT": ["T"] * n_rows,
            "GT": ["Sample1(0/1)"] * n_rows,
            "ClinVar": [f"https://example.org/cv/{i}" for i in range(n_rows)],
        }
    )

    extra_sheets = {
        "Metadata": pd.DataFrame({"K": ["version"], "V": ["v1.0"]}),
        "Statistics": pd.DataFrame({"Metric": ["count"], "Value": [n_rows]}),
    }

    # 1. Multi-pass reload baseline (historical pre-#159 sequence: convert -> append -> finalize)
    # Run each step sequentially with explicit workbook close and gc.collect
    f_baseline = tmp_path / "baseline.xlsx"
    gc.collect()
    tracemalloc.start()
    t0_base = time.perf_counter()

    # Step 0: Initial dataframe write
    df.to_excel(f_baseline, sheet_name="Results", index=False)

    # Step 1: Append Metadata sheet
    wb1 = openpyxl.load_workbook(f_baseline)
    ws1 = wb1.create_sheet("Metadata")
    ws1.append(["K", "V"])
    ws1.append(["version", "v1.0"])
    wb1.save(f_baseline)
    wb1.close()
    del wb1
    gc.collect()

    # Step 2: Append Statistics sheet
    wb2 = openpyxl.load_workbook(f_baseline)
    ws2 = wb2.create_sheet("Statistics")
    ws2.append(["Metric", "Value"])
    ws2.append(["count", n_rows])
    wb2.save(f_baseline)
    wb2.close()
    del wb2
    gc.collect()

    # Step 3: Format links, autofilter, freeze panes
    wb3 = openpyxl.load_workbook(f_baseline)
    for ws in wb3.worksheets:
        ws.freeze_panes = "A2"
        max_col_letter = get_column_letter(ws.max_column)
        ws.auto_filter.ref = f"A1:{max_col_letter}1"
        if ws.title == "Results":
            header_row = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            clinvar_idx = header_row.index("ClinVar") + 1
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=clinvar_idx)
                val = str(cell.value or "").strip()
                if val.startswith(("http://", "https://")):
                    cell.hyperlink = val
                    cell.value = "ClinVar"
                    cell.style = "Hyperlink"
    wb3.save(f_baseline)
    wb3.close()
    del wb3
    gc.collect()

    elapsed_base = time.perf_counter() - t0_base
    _current, peak_base = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    gc.collect()

    # 2. Single-pass xlsxwriter generator (#159)
    f_single = tmp_path / "single.xlsx"
    gc.collect()
    tracemalloc.start()
    t0_single = time.perf_counter()
    convert_to_excel(
        str(f_single),
        cfg={"igv_enabled": False},
        df=df,
        extra_sheets=extra_sheets,
        output_file=str(f_single),
    )
    elapsed_single = time.perf_counter() - t0_single
    _current, peak_single = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    gc.collect()

    peak_base_mb = peak_base / (1024 * 1024)
    peak_single_mb = peak_single / (1024 * 1024)

    record_property("baseline_elapsed_s", elapsed_base)
    record_property("single_pass_elapsed_s", elapsed_single)
    record_property("baseline_peak_mb", peak_base_mb)
    record_property("single_pass_peak_mb", peak_single_mb)

    # Sanity checks
    assert elapsed_base > 0 and elapsed_single > 0
    assert peak_single_mb < 50.0, f"Peak memory {peak_single_mb:.1f}MB, expected < 50MB"

    # Cell-by-cell value and hyperlink fidelity comparison across all sheets
    wb_single = openpyxl.load_workbook(f_single)
    wb_base = openpyxl.load_workbook(f_baseline)
    assert wb_single.sheetnames == wb_base.sheetnames

    for sheet_name in wb_single.sheetnames:
        ws_single = wb_single[sheet_name]
        ws_base = wb_base[sheet_name]
        assert ws_single.max_row == ws_base.max_row
        assert ws_single.max_column == ws_base.max_column

        for r in range(1, ws_single.max_row + 1):
            for c in range(1, ws_single.max_column + 1):
                cell_s = ws_single.cell(row=r, column=c)
                cell_b = ws_base.cell(row=r, column=c)
                assert cell_s.value == cell_b.value, (
                    f"Mismatch at {sheet_name} row {r} col {c}: {cell_s.value} != {cell_b.value}"
                )
                if cell_b.hyperlink:
                    assert cell_s.hyperlink is not None, (
                        f"Missing hyperlink at {sheet_name} row {r} col {c}"
                    )
                    assert cell_s.hyperlink.target == cell_b.hyperlink.target


@pytest.mark.unit
def test_excel_mixed_url_column_hyperlinks_preserved(tmp_path: Path):
    """Verify URLs in mixed columns (<70% URLs) are clickable hyperlinks (#159)."""
    from variantcentrifuge.converter import convert_to_excel

    df = pd.DataFrame(
        {
            "CHROM": ["chr1", "chr1", "chr1", "chr1"],
            "POS": [100, 200, 300, 400],
            "REF": ["A", "C", "G", "T"],
            "ALT": ["T", "G", "A", "C"],
            "MixedCol": [
                "unavailable",
                "unavailable",
                "https://example.org/a",
                "https://example.org/b",
            ],
        }
    )
    xlsx_file = tmp_path / "mixed.xlsx"
    convert_to_excel(
        str(xlsx_file),
        cfg={"igv_enabled": False},
        df=df,
        output_file=str(xlsx_file),
    )

    wb = openpyxl.load_workbook(xlsx_file)
    ws = wb["Results"]
    mixed_col_idx = 5

    # Row 2 & 3: plain text strings, no hyperlink
    cell_row2 = ws.cell(row=2, column=mixed_col_idx)
    assert cell_row2.value == "unavailable"
    assert cell_row2.hyperlink is None

    cell_row3 = ws.cell(row=3, column=mixed_col_idx)
    assert cell_row3.value == "unavailable"
    assert cell_row3.hyperlink is None

    # Row 4 & 5: clickable hyperlinks displaying the URL string
    cell_row4 = ws.cell(row=4, column=mixed_col_idx)
    assert cell_row4.value == "https://example.org/a"
    assert cell_row4.hyperlink is not None
    assert cell_row4.hyperlink.target == "https://example.org/a"

    cell_row5 = ws.cell(row=5, column=mixed_col_idx)
    assert cell_row5.value == "https://example.org/b"
    assert cell_row5.hyperlink is not None
    assert cell_row5.hyperlink.target == "https://example.org/b"


@pytest.mark.unit
def test_finalize_excel_file_overflow_sheet_created(tmp_path: Path):
    """Verify finalize_excel_file creates IGV_Overflow sheet on >32,000 char comments (#159)."""
    from variantcentrifuge.converter import finalize_excel_file

    output_dir = tmp_path / "finalize_overflow"
    output_dir.mkdir(parents=True, exist_ok=True)
    report_dir = output_dir / "report" / "igv"
    report_dir.mkdir(parents=True, exist_ok=True)

    n_samples = 600
    subfolder = "igv/long_batch_path_subfolder"
    sample_reports = {
        f"Patient_Sample_ID_{i:04d}": f"{subfolder}/Patient_Sample_ID_{i:04d}_variant_report.html"
        for i in range(n_samples)
    }
    igv_map = {
        "variants": [
            {
                "chrom": "chr1",
                "pos": 100,
                "ref": "A",
                "alt": "T",
                "sample_reports": sample_reports,
            }
        ]
    }
    with open(report_dir / "igv_reports_map.json", "w", encoding="utf-8") as f:
        json.dump(igv_map, f)

    df = pd.DataFrame(
        {
            "CHROM": ["chr1"],
            "POS": [100],
            "REF": ["A"],
            "ALT": ["T"],
            "GT": [";".join(f"Patient_Sample_ID_{i:04d}(0/1)" for i in range(n_samples))],
        }
    )
    xlsx_file = output_dir / "finalize_overflow.xlsx"
    df.to_excel(xlsx_file, sheet_name="Results", index=False)

    finalize_excel_file(str(xlsx_file), cfg={"igv_enabled": True})

    wb = openpyxl.load_workbook(xlsx_file)
    assert "Results" in wb.sheetnames
    assert "IGV_Overflow" in wb.sheetnames

    ws_overflow = wb["IGV_Overflow"]
    expected_headers = ["Row", "CHROM", "POS", "REF", "ALT", "Sample_ID", "IGV_Report_Path"]
    assert [c.value for c in ws_overflow[1]] == expected_headers
    assert ws_overflow.max_row == n_samples + 1

    first_row = [
        ws_overflow.cell(row=2, column=i).value for i in range(1, len(expected_headers) + 1)
    ]
    assert first_row == [
        2,
        "chr1",
        100,
        "A",
        "T",
        "Patient_Sample_ID_0000",
        "report/igv/long_batch_path_subfolder/Patient_Sample_ID_0000_variant_report.html",
    ]


@pytest.mark.unit
def test_finalize_excel_file_missing_coordinates_guard(tmp_path: Path):
    """Verify finalize_excel_file handles missing coordinates without column 0 ValueError (#159)."""
    from variantcentrifuge.converter import finalize_excel_file

    output_dir = tmp_path / "missing_coords"
    output_dir.mkdir(parents=True, exist_ok=True)
    report_dir = output_dir / "report" / "igv"
    report_dir.mkdir(parents=True, exist_ok=True)

    igv_map = {
        "variants": [
            {
                "chrom": "chr1",
                "pos": 100,
                "ref": "A",
                "alt": "T",
                "sample_reports": {"Sample1": "igv/Sample1.html"},
            }
        ]
    }
    with open(report_dir / "igv_reports_map.json", "w", encoding="utf-8") as f:
        json.dump(igv_map, f)

    # DataFrame with only CHROM and POS (missing REF and ALT)
    df = pd.DataFrame({"CHROM": ["chr1"], "POS": [100]})
    xlsx_file = output_dir / "partial.xlsx"
    df.to_excel(xlsx_file, sheet_name="Results", index=False)

    # Should safely skip IGV links rather than calling ws.cell(..., column=0)
    finalize_excel_file(str(xlsx_file), cfg={"igv_enabled": True})

    wb = openpyxl.load_workbook(xlsx_file)
    ws = wb["Results"]
    headers = [cell.value for cell in ws[1]]
    assert "IGV Report Links" not in headers
    assert headers == ["CHROM", "POS"]


@pytest.mark.unit
def test_excel_igv_comment_overflow_creates_overflow_sheet(tmp_path: Path):
    """Verify IGV comment >32,000 chars is truncated and full links saved to IGV_Overflow (#159)."""
    from variantcentrifuge.converter import convert_to_excel

    output_dir = tmp_path / "overflow_test"
    output_dir.mkdir(parents=True, exist_ok=True)
    report_dir = output_dir / "report" / "igv"
    report_dir.mkdir(parents=True, exist_ok=True)

    # 600 samples with long paths -> comment will exceed 32,000 characters
    n_samples = 600
    subfolder = "igv/long_batch_path_subfolder"
    sample_reports = {
        f"Patient_Sample_ID_{i:04d}": f"{subfolder}/Patient_Sample_ID_{i:04d}_variant_report.html"
        for i in range(n_samples)
    }
    igv_map = {
        "variants": [
            {
                "chrom": "chr1",
                "pos": 100,
                "ref": "A",
                "alt": "T",
                "sample_reports": sample_reports,
            }
        ]
    }
    with open(report_dir / "igv_reports_map.json", "w", encoding="utf-8") as f:
        json.dump(igv_map, f)

    df = pd.DataFrame(
        {
            "CHROM": ["chr1"],
            "POS": [100],
            "REF": ["A"],
            "ALT": ["T"],
            "GT": [";".join(f"Patient_Sample_ID_{i:04d}(0/1)" for i in range(n_samples))],
        }
    )
    xlsx_file = output_dir / "overflow.xlsx"
    convert_to_excel(
        str(xlsx_file),
        cfg={"igv_enabled": True},
        df=df,
        output_file=str(xlsx_file),
    )

    wb = openpyxl.load_workbook(xlsx_file)
    assert "Results" in wb.sheetnames
    assert "IGV_Overflow" in wb.sheetnames

    ws_results = wb["Results"]
    headers = [cell.value for cell in ws_results[1]]
    assert "IGV Report Links" in headers
    igv_col = headers.index("IGV Report Links") + 1

    cell = ws_results.cell(row=2, column=igv_col)
    assert cell.value == f"Patient_Sample_ID_0000 (+{n_samples - 1} others)"
    assert cell.comment is not None

    comment_text = cell.comment.text
    # Cell comment must be at most 32,767 chars (Excel limit) and end with overflow suffix
    assert len(comment_text) <= 32767
    assert len(comment_text) <= 32000
    expected_suffix = "\n... (see IGV_Overflow sheet for all sample links)"
    assert comment_text.endswith(expected_suffix)

    # Verify IGV_Overflow sheet
    ws_overflow = wb["IGV_Overflow"]
    overflow_headers = [c.value for c in ws_overflow[1]]
    expected_headers = ["Row", "CHROM", "POS", "REF", "ALT", "Sample_ID", "IGV_Report_Path"]
    assert overflow_headers == expected_headers

    # 1 header row + 600 sample rows
    assert ws_overflow.max_row == n_samples + 1
    first_row = [
        ws_overflow.cell(row=2, column=i).value for i in range(1, len(expected_headers) + 1)
    ]
    assert first_row == [
        2,
        "chr1",
        100,
        "A",
        "T",
        "Patient_Sample_ID_0000",
        "report/igv/long_batch_path_subfolder/Patient_Sample_ID_0000_variant_report.html",
    ]


@pytest.mark.unit
def test_produce_report_json_numeric_gt(tmp_path: Path):
    """Verify produce_report_json handles numeric GT (haploid/unquoted) without crashing (#159)."""
    from variantcentrifuge.converter import produce_report_json

    output_dir = tmp_path / "report_json_test"
    output_dir.mkdir(parents=True, exist_ok=True)
    report_dir = output_dir / "report" / "igv"
    report_dir.mkdir(parents=True, exist_ok=True)

    igv_map = {
        "variants": [
            {
                "chrom": "chr1",
                "pos": 100,
                "ref": "A",
                "alt": "T",
                "sample_reports": {"Sample1": "igv/Sample1.html", "1": "igv/sample_1.html"},
            }
        ]
    }
    with open(report_dir / "igv_reports_map.json", "w", encoding="utf-8") as f:
        json.dump(igv_map, f)

    tsv_path = output_dir / "variants.tsv"
    tsv_content = "CHROM\tPOS\tREF\tALT\tGT\nchr1\t100\tA\tT\t1\nchr1\t200\tG\tC\t0\n"
    tsv_path.write_text(tsv_content, encoding="utf-8")

    produce_report_json(str(tsv_path), str(output_dir))

    variants_json_path = output_dir / "report" / "variants.json"
    summary_json_path = output_dir / "report" / "summary.json"
    assert variants_json_path.exists()
    assert summary_json_path.exists()

    with open(variants_json_path, encoding="utf-8") as f:
        variants = json.load(f)

    assert len(variants) == 2
    # First variant (numeric GT=1) should be loaded without crashing
    assert variants[0]["GT"] == 1
    assert variants[0]["igv_links"] == []
    # Second variant (numeric GT=0) reference genotype -> no igv_links
    assert variants[1]["GT"] == 0
    assert variants[1]["igv_links"] == []

    # Verify _get_igv_reports_for_variant handles int, None, and haploid string directly
    from variantcentrifuge.converter import _get_igv_reports_for_variant

    igv_lookup = {
        ("chr1", "100", "A", "T", "Sample1"): "igv/Sample1.html",
    }
    assert _get_igv_reports_for_variant("chr1", "100", "A", "T", 1, igv_lookup) == []
    assert _get_igv_reports_for_variant("chr1", "100", "A", "T", 0, igv_lookup) == []
    assert _get_igv_reports_for_variant("chr1", "100", "A", "T", "Sample1(1)", igv_lookup) == [
        ("Sample1", "igv/Sample1.html")
    ]
    assert _get_igv_reports_for_variant("chr1", "100", "A", "T", "Sample1(0)", igv_lookup) == []
    assert _get_igv_reports_for_variant("chr1", "100", "A", "T", "Sample1(.)", igv_lookup) == []
    assert _get_igv_reports_for_variant("chr1", "100", "A", "T", None, igv_lookup) == []


@pytest.mark.unit
def test_excel_hyperlink_protocols_preserved(tmp_path: Path):
    """Verify ftp, ftps, mailto, internal, and external links are preserved in Excel (#159)."""
    from variantcentrifuge.converter import convert_to_excel

    output_dir = tmp_path / "hyperlinks_test"
    output_dir.mkdir(parents=True, exist_ok=True)
    xlsx_file = output_dir / "links.xlsx"

    df = pd.DataFrame(
        {
            "CHROM": ["chr1"],
            "POS": [100],
            "REF": ["A"],
            "ALT": ["T"],
            "FTP_LINK": ["ftp://ftp.ncbi.nlm.nih.gov/genomes/all/data.txt"],
            "FTPS_LINK": ["ftps://secure.example.org/path/file.bin"],
            "EMAIL": ["mailto:curator@genomics.org"],
            "INTERNAL_REF": ["internal:Sheet1!A1"],
            "EXTERNAL_REF": ["external:doc.xlsx#Sheet1!A1"],
            "HTTP_URL": ["https://www.ncbi.nlm.nih.gov/clinvar/"],
        }
    )

    convert_to_excel(
        str(xlsx_file),
        cfg={"igv_enabled": False},
        df=df,
        output_file=str(xlsx_file),
    )

    wb = openpyxl.load_workbook(xlsx_file)
    ws = wb["Results"]
    headers = [cell.value for cell in ws[1]]
    col_indices = {name: idx + 1 for idx, name in enumerate(headers)}

    # FTP link
    cell_ftp = ws.cell(row=2, column=col_indices["FTP_LINK"])
    assert cell_ftp.value == "ftp://ftp.ncbi.nlm.nih.gov/genomes/all/data.txt"
    assert cell_ftp.hyperlink is not None
    assert cell_ftp.hyperlink.target == "ftp://ftp.ncbi.nlm.nih.gov/genomes/all/data.txt"

    # FTPS link
    cell_ftps = ws.cell(row=2, column=col_indices["FTPS_LINK"])
    assert cell_ftps.value == "ftps://secure.example.org/path/file.bin"
    assert cell_ftps.hyperlink is not None
    assert cell_ftps.hyperlink.target == "ftps://secure.example.org/path/file.bin"

    # Mailto link: prefix stripped for display, hyperlink target keeps mailto:
    cell_mail = ws.cell(row=2, column=col_indices["EMAIL"])
    assert cell_mail.value == "curator@genomics.org"
    assert cell_mail.hyperlink is not None
    assert cell_mail.hyperlink.target == "mailto:curator@genomics.org"

    # Internal ref link: XlsxWriter native displays location
    cell_int = ws.cell(row=2, column=col_indices["INTERNAL_REF"])
    assert cell_int.value == "Sheet1!A1"
    assert cell_int.hyperlink is not None
    assert cell_int.hyperlink.location == "Sheet1!A1"

    # External ref link: XlsxWriter native strips external: prefix
    cell_ext = ws.cell(row=2, column=col_indices["EXTERNAL_REF"])
    assert cell_ext.value == "doc.xlsx#Sheet1!A1"
    assert cell_ext.hyperlink is not None
    assert cell_ext.hyperlink.target == "doc.xlsx"

    # HTTP link in URL column: displays column name
    cell_http = ws.cell(row=2, column=col_indices["HTTP_URL"])
    assert cell_http.value == "HTTP_URL"
    assert cell_http.hyperlink is not None
    assert cell_http.hyperlink.target == "https://www.ncbi.nlm.nih.gov/clinvar/"


@pytest.mark.unit
def test_float_integral_pos_normalized_for_igv_lookup(tmp_path: Path):
    """Verify float POS (e.g. 100.0) matches IGV lookup with POS='100' in Excel (#159)."""
    from variantcentrifuge.converter import convert_to_excel, finalize_excel_file

    output_dir = tmp_path / "float_pos_test"
    output_dir.mkdir(parents=True, exist_ok=True)
    report_dir = output_dir / "report" / "igv"
    report_dir.mkdir(parents=True, exist_ok=True)

    igv_map = {
        "variants": [
            {
                "chrom": "chr1",
                "pos": 100,
                "ref": "A",
                "alt": "T",
                "sample_reports": {"Sample1": "igv/Sample1.html"},
            }
        ]
    }
    with open(report_dir / "igv_reports_map.json", "w", encoding="utf-8") as f:
        json.dump(igv_map, f)

    df = pd.DataFrame(
        {
            "CHROM": ["chr1"],
            "POS": [100.0],
            "REF": ["A"],
            "ALT": ["T"],
            "GT": ["Sample1(0/1)"],
        }
    )
    assert df["POS"].dtype == "float64"

    # Test convert_to_excel
    xlsx_convert = output_dir / "convert.xlsx"
    convert_to_excel(
        str(xlsx_convert),
        cfg={"igv_enabled": True},
        df=df,
        output_file=str(xlsx_convert),
    )

    wb_convert = openpyxl.load_workbook(xlsx_convert)
    ws_convert = wb_convert["Results"]
    headers_convert = [c.value for c in ws_convert[1]]
    assert "IGV Report Links" in headers_convert
    igv_col_convert = headers_convert.index("IGV Report Links") + 1
    cell_convert = ws_convert.cell(row=2, column=igv_col_convert)
    assert cell_convert.value == "Sample1"
    assert cell_convert.hyperlink is not None
    assert cell_convert.hyperlink.target.replace("\\", "/") == "report/igv/Sample1.html"

    # Test finalize_excel_file
    xlsx_finalize = output_dir / "finalize.xlsx"
    df.to_excel(xlsx_finalize, sheet_name="Results", index=False)
    finalize_excel_file(str(xlsx_finalize), cfg={"igv_enabled": True})

    wb_finalize = openpyxl.load_workbook(xlsx_finalize)
    ws_finalize = wb_finalize["Results"]
    headers_finalize = [c.value for c in ws_finalize[1]]
    assert "IGV Report Links" in headers_finalize
    igv_col_finalize = headers_finalize.index("IGV Report Links") + 1
    cell_finalize = ws_finalize.cell(row=2, column=igv_col_finalize)
    assert cell_finalize.value == "Sample1"
    assert cell_finalize.hyperlink is not None
    assert cell_finalize.hyperlink.target.replace("\\", "/") == "report/igv/Sample1.html"


@pytest.mark.unit
def test_excel_url_sampling_leading_nulls_and_native_links(tmp_path: Path):
    """Verify leading nulls not skipped in 4-row URL check and native links preserved (#159)."""
    from variantcentrifuge.converter import _detect_url_columns, convert_to_excel

    # 1. Verify _detect_url_columns with leading nulls
    df_sampling = pd.DataFrame(
        {
            "LeadingNullsCol": [
                None,
                None,
                None,
                "plain",
                "https://example.org/1",
                "https://example.org/2",
                "https://example.org/3",
            ],
            "AllUrlCol": ["https://example.org/a"] * 7,
        }
    )
    detected = _detect_url_columns(df_sampling)
    assert "LeadingNullsCol" not in detected
    assert "AllUrlCol" in detected

    # 2. Verify convert_to_excel native link generation
    xlsx_file = tmp_path / "native_links.xlsx"
    df_links = pd.DataFrame(
        {
            "CHROM": ["chr1"] * 7,
            "POS": list(range(100, 107)),
            "REF": ["A"] * 7,
            "ALT": ["T"] * 7,
            "InternalCol": [f"internal:Sheet1!A{i}" for i in range(1, 8)],
            "ExternalCol": [f"external:doc.xlsx#Sheet1!A{i}" for i in range(1, 8)],
            "LeadingNullsCol": [
                None,
                None,
                None,
                "plain",
                "https://example.org/1",
                "https://example.org/2",
                "https://example.org/3",
            ],
            "UrlCol": [f"https://example.org/{i}" for i in range(1, 8)],
        }
    )

    convert_to_excel(
        str(xlsx_file),
        cfg={"igv_enabled": False},
        df=df_links,
        output_file=str(xlsx_file),
    )

    wb = openpyxl.load_workbook(xlsx_file)
    ws = wb["Results"]

    # Internal links
    c_int = ws.cell(row=2, column=5)
    assert c_int.value == "Sheet1!A1"
    assert c_int.hyperlink is not None
    assert c_int.hyperlink.location == "Sheet1!A1"

    # External links
    c_ext = ws.cell(row=2, column=6)
    assert c_ext.value == "doc.xlsx#Sheet1!A1"
    assert c_ext.hyperlink is not None
    assert c_ext.hyperlink.target == "doc.xlsx"

    # LeadingNullsCol row 5 (row_idx 3): plain text
    c_plain = ws.cell(row=5, column=7)
    assert c_plain.value == "plain"
    assert c_plain.hyperlink is None

    # LeadingNullsCol row 6 (row_idx 4, url 1): not in url_columns -> displays full URL
    c_nulls = ws.cell(row=6, column=7)
    assert c_nulls.value == "https://example.org/1"
    assert c_nulls.hyperlink is not None
    assert c_nulls.hyperlink.target == "https://example.org/1"

    # UrlCol row 2: in url_columns -> displays column name "UrlCol"
    c_url = ws.cell(row=2, column=8)
    assert c_url.value == "UrlCol"
    assert c_url.hyperlink is not None
    assert c_url.hyperlink.target == "https://example.org/1"
